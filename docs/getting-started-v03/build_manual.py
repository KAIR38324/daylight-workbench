"""Build the Chinese getting-started Word guide from a completed, real journey.

Use the bundled Codex Python runtime, after running the documents artifact marker.
Screenshots are embedded unchanged from docs/getting-started; no synthetic UI frames.
"""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
REPORT = ROOT / 'qa/evidence/user-journey-v03/result.json'
OUTPUT = ROOT / 'docs/建筑采光多目标优化平台 入门指南 v0.3.docx'
FONT = 'Noto Sans CJK SC'
DATA = json.loads(REPORT.read_text(encoding='utf-8'))


def font(style, size, bold=False, color=(0, 0, 0)):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(*color)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)


def text(doc, value, style=None, after=6):
    p = doc.add_paragraph(style=style)
    p.add_run(value)
    p.paragraph_format.space_after = Pt(after)
    return p


def cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def borders(table):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.first_child_found_in('w:tblBorders')
    if old is not None:
        tbl_pr.remove(old)
    box = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{edge}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:color'), 'D9D9D9')
        border.set(qn('w:sz'), '4')
        box.append(border)
    tbl_pr.append(box)


def set_padding(table, top=100, start=100, bottom=100, end=100):
    pr = table._tbl.tblPr
    margins = OxmlElement('w:tblCellMar')
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        el = OxmlElement(f'w:{name}')
        el.set(qn('w:w'), str(value))
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    pr.append(margins)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = title
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_shading(cell, 'EDF1F4')
        for p in cell.paragraphs:
            p.style = doc.styles['Table Text']
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.style = doc.styles['Table Text']
                p.paragraph_format.space_after = Pt(0)
    borders(table)
    set_padding(table)
    return table


def f(value, digits=3):
    return f'{value:.{digits}f}'


steps = [
    ('导入合成 CSV', f"操作：点击“导入 CSV / XLSX”，选择 demo/manual-synthetic.csv。检查：出现 {DATA['row_count']} 行、13 列；文件名包含 synthetic。", None),
    ('明确输入目标与方向', '操作：将前 10 列设为“自变量 X”，后三列设为“目标 Y”；采光分数最大化，眩光与能耗分数最小化。检查：每个目标均有方向，不靠列名猜测物理含义。', None),
    ('校验数据与划分', f"操作：点击“校验数据与划分”。检查：本次有效 240 行，训练 {DATA['split_report']['split_counts']['train']}、验证 {DATA['split_report']['split_counts']['validation']}、测试 {DATA['split_report']['split_counts']['test']} 行；重复输入和同组样本不跨集。", None),
    ('设置逐目标 ANN', '操作：切到“模型与训练”，保留默认逐目标独立模式，设置 64 ReLU、32 Tanh 两个隐藏层；最多 180 轮，早停耐心 24，学习率 0.003，批量 32。检查：每个目标训练一套独立标量模型，而非固定共享三输出网络。', None),
    ('观察单次 ANN 实时训练', f"操作：点击“单次 ANN（快速）”。检查：轮次和真实训练、验证损失随后台任务更新。本次截图发生在第 {next(s['epoch'] for s in DATA['screenshots'] if s['number']==5)} 轮，任务仍在运行。", None),
    ('检查逐目标 ANN 指标', '操作：训练结束后打开“指标 / 实时曲线”，用目标选择框切换诊断。检查：三套独立 ANN 的测试 MAE、MSE、R² 及各自真实训练记录。此为固定划分快速运行，不与后面的交叉验证排名混同。', None),
    ('配置逐目标模型比较', '操作：在“模型选择 / 网格 / CV”确认默认“逐目标独立建模”，勾选比较 ANN、RF 与 XGBoost，严格模式保留 5 折。设 ANN 学习率 0.003、批量 32；RF 40 棵树、深度 8；XGBoost 30 轮、深度 3、学习率 0.1。共享多输出仅是可选兼容模式。', None),
    ('预览预算后开始验证', f"操作：先点击“预览候选与拟合预算”，再点击“开始网格验证 / 模型比较”。检查：3 目标 × 3 候选 × 5 折，再加 3 次开发集重拟合，实际拟合预算 {DATA['search_summary']['fit_budget']}/50；锁定测试 {DATA['search_locked_test_rows']} 行。", None),
    ('检查逐目标选优与测试', f"操作：展开逐目标候选排名，切换目标诊断并查看测试指标。本次各目标胜出模型：" + '；'.join(f'{name} → {family.upper()}' for name, family in DATA['per_target_winners'].items()) + f"。即使模型族相同，也是按各自验证分数独立选出的。仅重拟合开发集 {DATA['search_development_rows']} 行后，才评估锁定的 {DATA['search_locked_test_rows']} 行。", None),
    ('确认变量搜索边界', '操作：在 NSGA-II 页逐列检查训练观测范围、可编辑上下限和三项目标方向；明确将 floor_level 设为整数。设置种群 32、代数 8、种子 42。边界盒本身不保证物理可行性。', None),
    ('观察 NSGA-II 实时前沿', f"操作：点击“运行 NSGA-II（仅代理预测）”。检查：后台任务期间显示实时代数、评估数与预测前沿；本次截图在第 {next(s['generation'] for s in DATA['screenshots'] if s['number']==11)} 代。", None),
    ('确认优化完成', f"检查：本次实际评估 {DATA['optimization_evaluations']} 个候选。它们全是代理模型给出的预测值，尚未经过 Radiance、EnergyPlus 或实测验证。", None),
    ('浏览二维与平行坐标', '操作：在结果页切换到“全部候选 / 外部数据行”。检查：本次 256 个候选显示在目标二维散点和平行坐标中，13 个坐标维度含 3 项目标和 10 项输入。Linux 离屏环境未验证 WebGL 三维视图。', None),
    ('用候选 ID 联动选择', f"操作：显示表格，在图中选 3 个方案，再按候选 ID 排序表格。检查：图表与表格仍指向同一组 ID {', '.join(map(str, DATA['selected_ids_before_filter']))}，不按旧行号错配。", None),
    ('查看目标两两关系', '操作：在交互分析中切换“目标矩阵”。检查：目标两两散点维持相同的已选候选高亮。', None),
    ('查看指标分布', '操作：切换“分布”视图。检查：已选方案叠加在各目标的候选分布中。', None),
    ('对照已选方案', '操作：切换“对比”视图。检查：按原始分数单位对照所选方案，不把多个目标压成未解释的综合分数。', None),
    ('阅读相关性热图', '操作：切换“相关性”视图。检查：这是当前候选输入与目标的相关关系；候选有选择偏差，不得把相关性解释为因果。', None),
    ('填写手动区间条件', f"操作：展开“区间筛选”，在预测采光分数下限输入 {next(s['lower_bound'] for s in DATA['screenshots'] if s['number']==19):.3f}。其他上下限留空。", None),
    ('应用筛选并核对行数', f"操作：点击“应用筛选”。检查：当前视图为 {DATA['filtered_rows']} 行，低于 256 行；筛选不改写预测值，也不重新定义全表的 Pareto 标记。", None),
    ('导出并核对选择', f"操作：选定候选 ID {', '.join(map(str, DATA['exported_ids']))}，点击“导出已选方案”。检查：CSV 与来源 JSON 一同生成；本次自动化逐列比对导出与原候选数值，并核对 ID 和 SHA-256。", None),
    ('保存 v3 项目', '操作：点击顶端“保存项目”，生成 .dwb。检查：状态栏显示已保存；v3 保存每个目标的独立模型、严格搜索来源与候选。项目未加密，不向不可信对象分享敏感数据。', None),
    ('重开 v3 项目', '操作：点击“打开项目”选择刚保存的可信 .dwb。检查：逐目标胜出模型与候选配置、优化候选、102 行筛选结果及两个已选 ID 恢复；二维散点和图表计数均重新高亮。', None),
    ('确认决策目标映射', '操作：切到“偏好决策 / DeepSeek”，按本合成数据定义确认三个目标的中文含义、优化方向和别名；单位留空，不从列名推断物理单位。点击“确认这些目标含义和方向”。', None),
    ('预览采光倾向请求', '操作：输入“帮我选更考虑采光的”，点击“预览发送内容”，检查 api.deepseek.com 的完整请求正文。真实使用时还需自行设置合法 API Key，发送可能产生费用。本次没有真实密钥，也不访问网络。', None),
    ('审核采光模拟建议', '测试说明：本页采用本地替代传输，模拟返回采光 60%、眩光 20%、能耗 20% 的结构化建议。真实使用需点击“发送”并审阅提供商返回内容。这里展示的是未应用建议，不声称 DeepSeek 在线验证成功。', None),
    ('显式应用采光权重', '操作：审核权重、范围、约束与固定全候选归一化后，点击“应用审核后的偏好并联动选择”。检查：右侧以本地计算的无量纲偏好损失列出现有候选，不新造预测值，也不宣称唯一物理最优。', None),
    ('查看偏好联动图', '操作：打开“交互图表”。检查：完整候选背景仍可见，推荐候选以原始稳定 ID 联动结果页的二维散点和平行坐标。当前 Linux 离屏环境只验证二维回退，未验证 WebGL 三维。', None),
    ('预览第二轮偏好', '操作：输入“我要能耗和眩光平衡的方案”，再次点击“预览发送内容”。检查：请求含上一轮已审核偏好，但不含候选 ID、训练数据行或候选明细行；仍须用户主动发送。', None),
    ('审核平衡模拟建议', '测试说明：第二轮同样由本地替代传输模拟，建议能耗与眩光各 50%、采光 0%。在点击应用前，上一轮实际推荐保持不变；模拟建议不等于在线服务结论。', None),
    ('应用平衡权重', '操作：点击“应用审核后的偏好并联动选择”。检查：新权重下本地重新排序、推荐与稳定 ID 联动；固定全候选归一化保证两轮对照口径一致，原始 Pareto 标记不被修改。', None),
    ('映射已有数值结果', '操作：点击“导入外部结果”，将示例 CSV 的 score_a、score_b 设为目标并指定最小化，不指定输入。检查：仅有目标列也能分析，系统不会虚构设计变量。', None),
    ('查看已有结果', '检查：源文件 3 行以稳定源行 ID 显示；目标仍叫 score_a、score_b。软件不会自动把外部数值认证为代理预测、仿真或实测。', None),
]
assert len(steps) == len(DATA['screenshots']) == 33

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(.74)
sec.top_margin, sec.bottom_margin = Inches(.72), Inches(.66)
for name, size, bold in [('Normal', 10.5, False), ('Title', 20, True), ('Heading 1', 14, True), ('Caption', 8.5, False)]:
    font(doc.styles[name], size, bold)
font(doc.styles.add_style('Table Text', 1), 9, False)
title_style_pr = doc.styles['Title']._element.get_or_add_pPr()
title_border = title_style_pr.find(qn('w:pBdr'))
if title_border is not None:
    title_style_pr.remove(title_border)
for name in ('Normal', 'Heading 1', 'Caption'):
    doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.styles[name].paragraph_format.line_spacing = 1.18

title = doc.add_paragraph(style='Title')
title.add_run('建筑采光多目标优化平台 入门指南')
text(doc, '使用可复现合成数据完成从导入、模型训练到结果导出的完整流程。', after=10)
text(doc, '本指南对应 v0.3.0 桌面源码界面。所示分数和优化候选仅用于演示软件操作，不是 Radiance、EnergyPlus、Ladybug Tools 或实测的物理验证结果。', after=10)
text(doc, f"界面源码版本：{DATA.get('source_commit') or 'v0.3.0 / 源码指纹 '+DATA['source_fingerprint_sha256'][:12]}。完整 SHA-256 指纹记录在结果 JSON；本次没有验证更新后的 Windows 安装包。", after=7)
text(doc, '操作顺序：数据与列定义 → 逐目标独立建模 → NSGA-II 优化 → 结果分析与导出 → 偏好决策 → 外部结果。截图来自同一次 Linux 离屏 Qt 自动化运行；文件选择器由测试路径代选，科学计算、后台任务、图表、导出、项目重开和本地偏好应用均执行应用实际路径。两轮 DeepSeek 响应由本地替代传输模拟，未访问在线 API。', after=10)
text(doc, '开始前准备', style='Heading 1')
text(doc, '使用项目 README 的安装方法启动源码，或在 Windows 版本可用时从完整分发目录启动 DaylightWorkbench.exe；不要只复制单个 EXE。本文操作使用 demo/manual-synthetic.csv：固定种子 20261020，240 行、10 个输入、3 个目标。', after=7)
text(doc, '当前 Windows 打包验收与下述 v0.3.0 源码截图不是同一次验证。离屏 Linux 也没有可用的 WebGL，因此下文只展示已实际运行的二维和平行坐标回退，不将三维标为通过。', after=9)
text(doc, '本次运行的可核对结果', style='Heading 1')
add_table(doc, ['项目', '本次实测'], [
    ('数据拆分', f"训练 {DATA['split_report']['split_counts']['train']} / 验证 {DATA['split_report']['split_counts']['validation']} / 测试 {DATA['split_report']['split_counts']['test']}"),
    ('逐目标严格搜索', f"3 目标 × 3 候选 × 5 折 + 3 重拟合，预算 {DATA['search_summary']['fit_budget']}/50"),
    ('逐目标选优', '；'.join(name.replace('synthetic_', '').replace('_score', '') + ' ' + family.upper() for name, family in DATA['per_target_winners'].items())),
    ('NSGA-II', f"{DATA['optimization_evaluations']} 次评估，{DATA['optimization_candidates']} 个候选"),
    ('手动筛选与导出', f"{DATA['filtered_rows']} 行；导出 ID {', '.join(map(str, DATA['exported_ids']))}"),
    ('偏好决策', '本地模拟 2 轮建议；均经人工审核后显式应用，在线 API 未调用'),
], [2.1, 4.8])

for index, ((heading, instruction, _), screenshot) in enumerate(zip(steps, DATA['screenshots']), start=1):
    assert screenshot['number'] == index
    heading_p = text(doc, f'步骤 {index:02d} {heading}', style='Heading 1', after=10)
    heading_p.paragraph_format.page_break_before = True
    text(doc, instruction, after=10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.add_run().add_picture(str(ROOT / screenshot['file']), width=Inches(6.95))
    caption = doc.add_paragraph(style='Caption')
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.add_run(f"图 {index:02d}  {screenshot['caption']}")
    if index in (9, 12, 14, 20, 21, 23, 26, 29, 31):
        note = {
            9: '测试集只用于所选模型的最终报告，不参与候选排名或早停。',
            12: '搜索边界内的“最优”只针对代理模型目标与本次定义的方向。',
            14: '排序或筛选后的联动以 candidate_id 为准。',
            20: '如需恢复完整结果，可清空条件后重新应用。',
            21: '导出的来源 JSON 可用于核对数据与筛选来源，不等于物理有效性认证。',
            23: '仅打开可信来源的 .dwb；项目内容未加密。',
            26: '本图虽显示应用接收建议的界面，但建议来自测试脚本中的本地替代传输，不是在线 DeepSeek 返回。',
            29: '预览不等于发送；若要真实联网，用户须自行配置密钥，并明确点击“发送”。',
            31: '偏好权重只是帮助比较当前候选，不改变代理预测值或物理可行性。',
        }[index]
        text(doc, note, after=4)
    if index == 9:
        metrics = []
        for metric in DATA['heldout_metrics']:
            name = metric['目标'].replace('synthetic_', '').replace('_score', '')
            metrics.append(f"{name}：MAE {f(metric['MAE'])}，MSE {f(metric['MSE'])}，R² {f(metric['R2'])}")
        text(doc, '独立测试集逐目标指标（原分数单位）：' + '；'.join(metrics) + '。', after=4)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
